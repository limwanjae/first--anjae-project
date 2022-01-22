from openpyxl import load_workbook
wb = load_workbook("sample.xlsx")
ws = wb.active

from openpyxl.chart import BarChart, Reference,LineChart
# b1:c11까지의 데이터를 차트로~
line_value = Reference(ws,min_row=1, max_row=11,min_col=2, max_col=3)
line_chart = LineChart()# 차트 종류를 설정(Bar, Line, Pie, ...)
line_chart.add_data(line_value, titles_from_data=True)# 차트 데이터 추가, 계열부분(영어, 수학) 추가
line_chart.title = "성적표" # 제목
line_chart.style = 10 #미리 정의된 스타일을 적용, 사용자가 개별 지정도 가능
line_chart.y_axis.title = "점수" # y축의 제목
line_chart.x_axis.title = "번호" # x축의 제목

ws.add_chart(line_chart, "E1")# 차트 넣을 위치 정의
wb.save("sample_line_chart.xlsx")

# 구글에서 openpyxl  검색을 하면 찾아보면 여러유형의 차트가 있음